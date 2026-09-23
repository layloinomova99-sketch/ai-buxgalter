"""Knowledge base — Creator uploads files (via Telegram/web), AI extracts summary + rules."""
import os
import io
import json
import base64
import logging
from pathlib import Path
from typing import Optional
from bson import ObjectId
from services import llm

from core.db import get_db, iso_now

log = logging.getLogger("knowledge")

import tempfile

TEXT_EXT = {".py", ".txt", ".md", ".json", ".csv", ".xml", ".yaml", ".yml", ".js", ".ts", ".sql", ".ini", ".cfg", ".log", ".html", ".bsl", ".os"}
IMAGE_EXT = {".jpg", ".jpeg", ".png", ".webp"}
MAX_CHARS = 120_000

EXTRACT_PROMPT = """Sen — AI BUXGALTER bilim bazasi tahlilchisisan. Creator senga hujjat yukladi.
Vazifang: hujjat mazmunini tahlil qilib, QAT'IY JSON qaytarish (boshqa matn yo'q):
{
  "title": "qisqa sarlavha",
  "summary": "3-6 jumlada o'zbek tilida hujjat nima haqida, qanday ma'lumot va logika bor",
  "key_facts": ["muhim fakt yoki tafsilot", "..."],
  "rules": [
    {"category": "one_c_rules|soliq_rules|eimzo_rules|integration_rules|product_mapping_rules|mxik_rules|receipt_rules|marketplace_rules|accounting_rules|company_rules|creator_instructions",
     "instruction": "AI uchun aniq, bajariladigan qoida (o'zbek tilida)",
     "priority": 50}
  ]
}
Qoidalar faqat hujjatda HAQIQATAN mavjud logikadan olinsin — o'zingdan o'ylab topma. Agar kod bo'lsa: API endpointlar, autentifikatsiya usuli, hujjat/maydon nomlari, hisob ustunlari, mapping logikasi kabi aniq texnik faktlarni key_facts va rules'ga kirit. Maksimum 15 qoida."""


def extract_text(path: Path) -> tuple[str, str]:
    """Returns (kind, text). kind: text|image|binary."""
    ext = path.suffix.lower()
    if ext in IMAGE_EXT:
        return "image", ""
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        return "text", "\n".join((p.extract_text() or "") for p in reader.pages)
    if ext == ".docx":
        import docx
        d = docx.Document(str(path))
        return "text", "\n".join(p.text for p in d.paragraphs)
    if ext in {".xlsx", ".xlsm"}:
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            out.append(f"## Sheet: {ws.title}")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i > 500: break
                out.append("\t".join("" if v is None else str(v) for v in row))
        return "text", "\n".join(out)
    if ext in TEXT_EXT or ext == "":
        try:
            return "text", path.read_text(encoding="utf-8", errors="ignore")
        except Exception:
            return "binary", ""
    try:
        return "text", path.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        return "binary", ""


def _parse_json(raw: str) -> dict:
    raw = raw.strip()
    if raw.startswith("```"):
        raw = raw.strip("`")
        if raw.lower().startswith("json"):
            raw = raw[4:]
    start, end = raw.find("{"), raw.rfind("}")
    if start == -1 or end == -1:
        return {"title": "", "summary": raw[:1500], "key_facts": [], "rules": []}
    try:
        return json.loads(raw[start:end + 1])
    except Exception:
        return {"title": "", "summary": raw[:1500], "key_facts": [], "rules": []}


async def analyze_with_ai(kind: str, text: str, path: Path, filename: str, note: Optional[str]) -> dict:
    intro = f"Fayl nomi: {filename}\n"
    if note:
        intro += f"Creator izohi: {note}\n"
    if kind == "image":
        b64 = base64.b64encode(path.read_bytes()).decode()
        ext = path.suffix.lower().lstrip(".")
        media = {"jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png", "webp": "image/webp"}.get(ext, "image/jpeg")
        content = [{"type": "image", "source": {"type": "base64", "media_type": media, "data": b64}},
                   {"type": "text", "text": intro + "Rasmdagi ma'lumotni tahlil qil."}]
    else:
        body = text[:MAX_CHARS]
        if len(text) > MAX_CHARS:
            body += f"\n\n[... qisqartirildi, jami {len(text)} belgi]"
        content = intro + "\n--- HUJJAT ---\n" + body
    resp = await llm.create(EXTRACT_PROMPT, [{"role": "user", "content": content}])
    return _parse_json(llm.text_of(resp))


async def ingest_file(user_id: str, filename: str, data: bytes, source: str = "telegram",
                      company_id: Optional[str] = None, note: Optional[str] = None) -> dict:
    db = get_db()
    doc = {
        "creator_id": user_id, "company_id": company_id, "filename": filename,
        "size": len(data), "source": source, "note": note, "status": "processing",
        "title": "", "summary": "", "key_facts": [], "rule_ids": [], "text_chars": 0,
        "created_at": iso_now(), "updated_at": iso_now(),
    }
    res = await db.knowledge_docs.insert_one(doc)
    doc_id = str(res.inserted_id)
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=Path(filename).suffix.lower())
    tmp.write(data); tmp.close()
    path = Path(tmp.name)

    try:
        kind, text = extract_text(path)
        if kind == "binary":
            raise ValueError("Bu fayl turini o'qib bo'lmadi")
        analysis = await analyze_with_ai(kind, text, path, filename, note)
        rule_ids = []
        for r in (analysis.get("rules") or [])[:15]:
            if not r.get("instruction"):
                continue
            rr = await db.ai_rules.insert_one({
                "creator_id": user_id, "company_id": company_id,
                "category": r.get("category") or "integration_rules",
                "instruction": r["instruction"],
                "structured_rule": {"knowledge_doc_id": doc_id, "filename": filename},
                "priority": int(r.get("priority") or 50), "status": "active",
                "source": "document", "version": 1,
                "created_at": iso_now(), "updated_at": iso_now(),
            })
            rule_ids.append(str(rr.inserted_id))
        await db.knowledge_docs.update_one({"_id": res.inserted_id}, {"$set": {
            "status": "ready", "title": analysis.get("title") or filename,
            "summary": analysis.get("summary") or "", "key_facts": analysis.get("key_facts") or [],
            "rule_ids": rule_ids, "text_chars": len(text), "text_excerpt": text[:20000], "updated_at": iso_now(),
        }})
        await db.audit_logs.insert_one({"who": user_id, "what": f"Bilim bazasiga fayl yuklandi: {filename} ({len(rule_ids)} qoida)",
                                        "when": iso_now(), "company_id": company_id, "source": source,
                                        "created_at": iso_now(), "updated_at": iso_now()})
        return {"id": doc_id, "status": "ready", "title": analysis.get("title") or filename,
                "summary": analysis.get("summary") or "", "key_facts": analysis.get("key_facts") or [],
                "rules_count": len(rule_ids)}
    except Exception as e:
        log.exception(f"ingest error: {e}")
        await db.knowledge_docs.update_one({"_id": res.inserted_id}, {"$set": {"status": "failed", "summary": str(e), "updated_at": iso_now()}})
        return {"id": doc_id, "status": "failed", "error": str(e)}
    finally:
        if path.exists():
            path.unlink()


def doc_public(d: dict) -> dict:
    return {
        "id": str(d["_id"]), "filename": d.get("filename"), "title": d.get("title"), "summary": d.get("summary"),
        "key_facts": d.get("key_facts") or [], "rules_count": len(d.get("rule_ids") or []), "status": d.get("status"),
        "source": d.get("source"), "company_id": d.get("company_id"), "size": d.get("size"), "note": d.get("note"),
        "created_at": d.get("created_at"),
    }


async def delete_doc(doc_id: str) -> bool:
    db = get_db()
    d = await db.knowledge_docs.find_one({"_id": ObjectId(doc_id)})
    if not d:
        return False
    ids = [ObjectId(i) for i in d.get("rule_ids") or []]
    if ids:
        await db.ai_rules.delete_many({"_id": {"$in": ids}})
    await db.knowledge_docs.delete_one({"_id": d["_id"]})
    return True


async def search_knowledge(query: str, company_id: Optional[str], limit: int = 5) -> list:
    db = get_db()
    q = {"status": "ready", "$or": [
        {"title": {"$regex": query, "$options": "i"}},
        {"summary": {"$regex": query, "$options": "i"}},
        {"key_facts": {"$elemMatch": {"$regex": query, "$options": "i"}}},
        {"filename": {"$regex": query, "$options": "i"}},
        {"text_excerpt": {"$regex": query, "$options": "i"}},
    ]}
    rows = []
    async for d in db.knowledge_docs.find(q).sort("_id", -1).limit(limit):
        if d.get("company_id") and company_id and d["company_id"] != company_id:
            continue
        rows.append({"id": str(d["_id"]), "title": d.get("title"), "summary": d.get("summary"), "key_facts": (d.get("key_facts") or [])[:10]})
    return rows
